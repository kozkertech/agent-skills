#!/usr/bin/env python3
"""
generate_xlsx.py — Generates a branded XLSX spreadsheet using openpyxl.

Creates a cover sheet and data sheets for price lists, reports,
or any document type that maps naturally to tabular data.
"""

from pathlib import Path
from typing import List
from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side
)
from openpyxl.utils import get_column_letter
import re


def generate_xlsx(doc_data: dict, output_dir: Path, slug: str) -> str:
    """Generate a branded XLSX file. Returns the output file path."""
    brand     = doc_data["brand"]
    template  = doc_data["template"]
    instance  = doc_data["instance_inputs"]
    content   = doc_data["content"]
    context   = doc_data["context"]
    skill_dir = Path(doc_data.get("skill_dir", "."))

    wb = Workbook()
    wb.remove(wb.active)  # Remove default empty sheet

    colours   = brand.get("colours", {})
    fonts_cfg = brand.get("fonts", {})

    primary_hex   = colours.get("primary",    "#1F3864").lstrip("#")
    secondary_hex = colours.get("secondary",  "#FF6E30").lstrip("#")
    # muted = mid-grey for text (captions, dates, footnotes)
    muted_hex     = colours.get("muted",      "#888888").lstrip("#")
    # surface = light tint for alternating row fills / card backgrounds
    surface_hex   = colours.get("surface",    "#F5F5F5").lstrip("#")
    text_hex      = colours.get("text",       "#1E1E1E").lstrip("#")
    heading_font  = fonts_cfg.get("heading", "Calibri")
    body_font     = fonts_cfg.get("body",    "Calibri")

    # Cover sheet
    _add_cover_sheet(wb, brand, template, instance,
                     primary_hex, secondary_hex, muted_hex,
                     heading_font, body_font, skill_dir)

    # Determine what data sheets to create based on document type and context
    doc_type = template.get("document_type", "price_list")

    if doc_type == "price_list" or "pricing.current" in context:
        _add_pricing_sheet(wb, context, brand,
                           primary_hex, secondary_hex, surface_hex, text_hex,
                           heading_font, body_font)

    if "services.catalog" in context:
        _add_services_sheet(wb, context, brand,
                            primary_hex, secondary_hex, surface_hex, text_hex,
                            heading_font, body_font)

    # Additional sheets from pre-authored content
    for section_id, section_content in content.items():
        if section_id in ("cover", "pricing", "services"):
            continue
        if section_content and len(section_content) > 50:
            _add_text_sheet(wb, section_id, section_content, brand,
                            primary_hex, secondary_hex, heading_font, body_font)

    out_path = output_dir / f"{slug}.xlsx"
    wb.save(str(out_path))
    return str(out_path)


# ── Sheets ────────────────────────────────────────────────────────────────────

def _add_cover_sheet(wb: Workbook, brand: dict, template: dict, instance: dict,
                     primary_hex: str, secondary_hex: str, muted_hex: str,
                     heading_font: str, body_font: str, skill_dir: Path):
    ws = wb.create_sheet("Cover", 0)
    ws.sheet_view.showGridLines = False

    # Column widths: A=accent stripe, B=main content, C=secondary content (logo anchor)
    ws.column_dimensions["A"].width = 2   # thin accent column
    ws.column_dimensions["B"].width = 42
    ws.column_dimensions["C"].width = 24

    brand_name     = brand.get("brand_name", "")
    doc_type_label = template.get("name", template.get("document_type", "Document").replace("_", " ").title())
    client         = instance.get("client_name", "")
    project        = instance.get("project_name", "")

    # Accent column (A) — solid secondary stripe for the full cover
    for row in range(1, 40):
        cell = ws.cell(row=row, column=1)
        cell.fill = PatternFill("solid", fgColor=secondary_hex)

    # Row 1-2: padding
    ws.row_dimensions[1].height = 8
    ws.row_dimensions[2].height = 8

    # Row 3: brand name — primary colour (readable on white bg)
    ws.row_dimensions[3].height = 28
    c = ws.cell(row=3, column=2, value=brand_name.upper())
    c.font      = Font(name=heading_font, size=11, bold=True, color=primary_hex)
    c.alignment = Alignment(vertical="center")

    # Row 4: thin accent divider
    ws.row_dimensions[4].height = 3
    for col in range(2, 8):
        ws.cell(row=4, column=col).fill = PatternFill("solid", fgColor=secondary_hex)

    # Row 5: padding
    ws.row_dimensions[5].height = 12

    # Row 6: document title — large, primary colour
    ws.row_dimensions[6].height = 40
    c = ws.cell(row=6, column=2, value=doc_type_label)
    c.font      = Font(name=heading_font, size=24, bold=True, color=primary_hex)
    c.alignment = Alignment(vertical="center")

    # Row 7: padding
    ws.row_dimensions[7].height = 8

    # Row 8: client / project subtitle
    subtitle = " · ".join(filter(None, [f"For {client}" if client else "", project]))
    if subtitle:
        ws.row_dimensions[8].height = 24
        c = ws.cell(row=8, column=2, value=subtitle)
        c.font      = Font(name=body_font, size=13, color="888888")
        c.alignment = Alignment(vertical="center")

    # Row 9: padding
    ws.row_dimensions[9].height = 8

    # Row 10: date — use muted_hex for caption text
    from datetime import date
    ws.row_dimensions[10].height = 18
    c = ws.cell(row=10, column=2, value=date.today().strftime("%B %d, %Y"))
    c.font      = Font(name=body_font, size=10, color=muted_hex)
    c.alignment = Alignment(vertical="center")

    # Logo — anchor to cell C3 (top-right of content area)
    _add_logo_to_cover(ws, brand, skill_dir, anchor_cell="C3")


def _add_pricing_sheet(wb: Workbook, context: dict, brand: dict,
                       primary_hex: str, secondary_hex: str, surface_hex: str,
                       text_hex: str, heading_font: str, body_font: str):
    ws = wb.create_sheet("Pricing")
    ws.sheet_view.showGridLines = False

    pricing_data  = context.get("pricing.current", {})
    services_data = context.get("services.catalog", {})

    # Build a lookup for service names + descriptions
    service_lookup: dict = {}
    for svc in services_data.get("services", []):
        service_lookup[svc.get("id", "")] = svc

    headers    = ["Service",    "Description",  "Price",   "Currency", "Billing",    "Valid From"]
    col_widths = [28,            40,             14,        10,          14,           14]

    _write_sheet_header(ws, "Pricing", headers, col_widths,
                        primary_hex, secondary_hex, surface_hex,
                        heading_font, body_font)

    prices    = pricing_data.get("prices", [])
    data_start = 3  # first data row
    for i, price_entry in enumerate(prices, start=data_start):
        service_id  = price_entry.get("service_id", "")
        svc_info    = service_lookup.get(service_id, {})
        name        = svc_info.get("name", service_id)
        description = svc_info.get("description", "")

        row_data = [
            name,
            description,
            price_entry.get("price", ""),
            price_entry.get("currency", "USD"),
            price_entry.get("billing",  ""),
            price_entry.get("valid_from", ""),
        ]

        row_fill = surface_hex if i % 2 == 0 else "FFFFFF"
        for j, value in enumerate(row_data, start=1):
            cell = ws.cell(row=i, column=j, value=value)
            cell.font      = Font(name=body_font, size=11, color=text_hex)
            cell.fill      = PatternFill("solid", fgColor=row_fill)
            cell.alignment = Alignment(vertical="center", wrap_text=(j == 2))
            ws.row_dimensions[i].height = 22
            if j == 3 and isinstance(value, (int, float)):
                cell.number_format = '#,##0.00'

    # Total row — SUM formula for the Price column
    total_row = data_start + len(prices)
    if len(prices) > 0:
        ws.row_dimensions[total_row].height = 26
        tc_label = ws.cell(row=total_row, column=1, value="Total")
        tc_label.font      = Font(name=body_font, size=11, bold=True, color=primary_hex)
        tc_label.fill      = PatternFill("solid", fgColor=surface_hex)
        tc_label.alignment = Alignment(vertical="center")

        price_col = "C"
        formula_cell = ws.cell(row=total_row, column=3,
                               value=f"=SUM({price_col}{data_start}:{price_col}{total_row-1})")
        formula_cell.font          = Font(name=body_font, size=11, bold=True, color=primary_hex)
        formula_cell.fill          = PatternFill("solid", fgColor=surface_hex)
        formula_cell.number_format = '#,##0.00'
        formula_cell.alignment     = Alignment(vertical="center")

        # Fill remaining cells in total row
        for j in range(2, len(headers) + 1):
            if j == 3:
                continue
            c = ws.cell(row=total_row, column=j)
            c.fill = PatternFill("solid", fgColor=surface_hex)

    # Notes row
    notes = pricing_data.get("notes", "")
    if notes:
        notes_row = total_row + 2
        ws.row_dimensions[notes_row].height = 18
        c = ws.cell(row=notes_row, column=1, value=f"Note: {notes}")
        c.font = Font(name=body_font, size=10, italic=True, color="888888")
        ws.merge_cells(start_row=notes_row, start_column=1,
                       end_row=notes_row, end_column=len(headers))
    else:
        notes_row = total_row + 1

    # Print settings
    ws.freeze_panes = "A3"
    ws.print_area   = f"A1:{get_column_letter(len(headers))}{notes_row}"
    ws.page_setup.orientation  = "landscape"
    ws.page_setup.fitToWidth   = 1
    ws.page_setup.fitToHeight  = 0
    ws.page_setup.fitToPage    = True
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _add_services_sheet(wb: Workbook, context: dict, brand: dict,
                        primary_hex: str, secondary_hex: str, surface_hex: str,
                        text_hex: str, heading_font: str, body_font: str):
    ws = wb.create_sheet("Services")
    ws.sheet_view.showGridLines = False

    services_data = context.get("services.catalog", {})
    headers    = ["Service",   "Description",   "Deliverables",  "Base Price", "Billing"]
    col_widths = [28,           45,               45,              14,            14]

    _write_sheet_header(ws, "Services", headers, col_widths,
                        primary_hex, secondary_hex, surface_hex,
                        heading_font, body_font)

    services  = services_data.get("services", [])
    last_data = 2
    for i, svc in enumerate(services, start=3):
        deliverables = ", ".join(svc.get("deliverables", []))
        row_data = [
            svc.get("name", ""),
            svc.get("description", ""),
            deliverables,
            svc.get("base_price", ""),
            svc.get("billing", ""),
        ]
        row_fill = surface_hex if i % 2 == 0 else "FFFFFF"
        for j, value in enumerate(row_data, start=1):
            cell = ws.cell(row=i, column=j, value=value)
            cell.font      = Font(name=body_font, size=11, color=text_hex)
            cell.fill      = PatternFill("solid", fgColor=row_fill)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            ws.row_dimensions[i].height = 40
            if j == 4 and isinstance(value, (int, float)):
                cell.number_format = '#,##0.00'
        last_data = i

    ws.freeze_panes = "A3"
    ws.print_area   = f"A1:{get_column_letter(len(headers))}{last_data}"
    ws.page_setup.orientation  = "landscape"
    ws.page_setup.fitToWidth   = 1
    ws.page_setup.fitToHeight  = 0
    ws.page_setup.fitToPage    = True
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _add_text_sheet(wb: Workbook, section_id: str, content_text: str, brand: dict,
                    primary_hex: str, secondary_hex: str,
                    heading_font: str, body_font: str):
    title = section_id.replace("_", " ").title()
    ws    = wb.create_sheet(title[:31])  # Sheet names max 31 chars
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 80

    # Heading row
    ws.row_dimensions[1].height = 28
    c = ws.cell(row=1, column=1, value=title)
    c.font      = Font(name=heading_font, size=14, bold=True, color=primary_hex)
    c.alignment = Alignment(vertical="center")

    # Accent row
    ws.row_dimensions[2].height = 3
    ws.cell(row=2, column=1).fill = PatternFill("solid", fgColor=secondary_hex)

    # Content paragraphs
    current_row = 3
    for para in re.split(r"\n\n+", content_text.strip()):
        para = para.strip()
        if not para:
            continue
        # Strip simple markdown markers for XLSX
        clean = re.sub(r"\*\*([^*]+)\*\*", r"\1", para)
        clean = re.sub(r"\*([^*]+)\*",     r"\1", clean)
        ws.row_dimensions[current_row].height = max(20, min(60, len(clean) // 4))
        c = ws.cell(row=current_row, column=1, value=clean)
        c.font      = Font(name=body_font, size=11, color="1E1E1E")
        c.alignment = Alignment(wrap_text=True, vertical="top")
        current_row += 1


# ── Logo ──────────────────────────────────────────────────────────────────────

def _add_logo_to_cover(ws, brand: dict, skill_dir: Path,
                       anchor_cell: str = "C3") -> bool:
    """
    Add the brand logo to the cover sheet using openpyxl's add_image().
    Returns True if placed, False if skipped.
    """
    logo_cfg = brand.get("logo", {})
    if not logo_cfg.get("enabled", False):
        return False

    logo_path_str = logo_cfg.get("path", "")
    if not logo_path_str:
        return False

    # Resolve path
    logo_file = None
    for candidate in [Path(logo_path_str), skill_dir / logo_path_str]:
        if candidate.exists():
            logo_file = candidate
            break

    if logo_file is None:
        return False

    try:
        from openpyxl.drawing.image import Image as XLImage
        img = XLImage(str(logo_file))
        # Scale to roughly 1.1 inches wide (openpyxl uses pixels at 96 DPI)
        xlsx_rules   = brand.get("logo_rules", {}).get("xlsx", {})
        width_inches = xlsx_rules.get("width_inches", 1.1)
        target_px    = int(width_inches * 96)
        if img.width and img.height:
            ratio     = img.height / img.width
            img.width  = target_px
            img.height = int(target_px * ratio)
        ws.add_image(img, anchor_cell)
        return True
    except Exception:
        return False


# ── Shared helpers ────────────────────────────────────────────────────────────

def _write_sheet_header(ws, title_text: str, headers: List[str], col_widths: List[int],
                        primary_hex: str, secondary_hex: str, surface_hex: str,
                        heading_font: str, body_font: str):
    """Write a branded title row and a styled column header row."""
    num_cols = len(headers)

    # Row 1: sheet title (merged across all columns)
    ws.row_dimensions[1].height = 30
    c = ws.cell(row=1, column=1, value=title_text)
    c.font      = Font(name=heading_font, size=14, bold=True, color=primary_hex)
    c.alignment = Alignment(vertical="center")
    c.fill      = PatternFill("solid", fgColor=surface_hex)
    if num_cols > 1:
        ws.merge_cells(start_row=1, start_column=1,
                       end_row=1,   end_column=num_cols)

    # Row 2: column headers
    ws.row_dimensions[2].height = 26
    for j, (header, width) in enumerate(zip(headers, col_widths), start=1):
        ws.column_dimensions[get_column_letter(j)].width = width
        cell = ws.cell(row=2, column=j, value=header)
        cell.font      = Font(name=heading_font, size=11, bold=True, color="FFFFFF")
        cell.fill      = PatternFill("solid", fgColor=primary_hex)
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border    = Border(
            bottom=Side(style="thin", color=secondary_hex)
        )
