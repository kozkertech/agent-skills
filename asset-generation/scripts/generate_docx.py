#!/usr/bin/env python3
"""
generate_docx.py — Generates a branded DOCX document using python-docx.

Receives a doc_data bundle containing plan, brand, template, context, instance_inputs,
pre-authored section content, and skill_dir. Applies brand styling and populates
sections in template order. Supports inline markdown (bold/italic), bullet lists,
and Markdown-style tables (| col | col |).
"""

from pathlib import Path
from typing import Optional, List
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import json
import re


def generate_docx(doc_data: dict, output_dir: Path, slug: str) -> str:
    """Generate a branded DOCX file. Returns the output file path."""
    brand     = doc_data["brand"]
    template  = doc_data["template"]
    instance  = doc_data["instance_inputs"]
    content   = doc_data["content"]
    context   = doc_data["context"]
    skill_dir = Path(doc_data.get("skill_dir", "."))

    doc = Document()
    _apply_page_setup(doc, brand)
    _add_logo_to_header(doc, brand, skill_dir)

    # Cover page
    doc_type = template.get("document_type", "")
    has_cover = brand.get("document_rules", {}).get("cover_page", True)
    if has_cover:
        _add_cover_page(doc, brand, template, instance)
        doc.add_page_break()

    # Table of Contents — only for multi-section documents (proposals, reports, contracts)
    toc_doc_types = {"proposal", "contract", "technical_report", "internal_report", "case_study"}
    if doc_type in toc_doc_types:
        _add_table_of_contents(doc, brand)

    # Body sections
    sections_order = template.get("required_sections", []) + [
        s for s in template.get("optional_sections", [])
        if s in (doc_data["plan"].get("include_optional_sections") or [])
    ]

    body_sections = [s for s in sections_order if s != "cover_page"]
    for idx, section_id in enumerate(body_sections):
        # Page break before every section except the very first (TOC already separates it)
        add_page_break = idx > 0
        _add_section(doc, section_id, content, context, instance, brand, template,
                     add_page_break=add_page_break)

    _add_footer(doc, brand)

    out_path = output_dir / f"{slug}.docx"
    doc.save(str(out_path))
    return str(out_path)


# ── Page setup ──────────────────────────────────────────────────────────────

def _apply_page_setup(doc: Document, brand: dict):
    spacing = brand.get("spacing", {})
    margin  = spacing.get("page_margin_inches", 1.0)
    colours = brand.get("colours", {})
    fonts   = brand.get("fonts", {})

    for section in doc.sections:
        section.top_margin    = Inches(margin)
        section.bottom_margin = Inches(margin)
        section.left_margin   = Inches(margin)
        section.right_margin  = Inches(margin)

    # Normal / body style
    style = doc.styles["Normal"]
    font  = style.font
    font.name       = fonts.get("body", "Calibri")
    font.size       = Pt(11)
    font.color.rgb  = _hex_to_rgb(colours.get("text", "#1E1E1E"))
    pf = style.paragraph_format
    pf.space_after       = Pt(8)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing      = Pt(11 * 1.15)

    # Heading 1 — section titles (enables nav pane, TOC, PDF bookmarks)
    primary_hex  = colours.get("primary", "#1F3864")
    heading_font = fonts.get("heading", "Calibri")
    try:
        h1 = doc.styles["Heading 1"]
        h1.font.name      = heading_font
        h1.font.size      = Pt(14)
        h1.font.bold      = True
        h1.font.color.rgb = _hex_to_rgb(primary_hex)
        h1.paragraph_format.space_before   = Pt(spacing.get("section_spacing_pt", 24))
        h1.paragraph_format.space_after    = Pt(2)
        h1.paragraph_format.keep_with_next = True
    except KeyError:
        pass

    # Heading 2 — sub-headings within sections (bold lines like **Phase 1 —**)
    try:
        h2 = doc.styles["Heading 2"]
        h2.font.name      = heading_font
        h2.font.size      = Pt(12)
        h2.font.bold      = True
        h2.font.color.rgb = _hex_to_rgb(primary_hex)
        h2.paragraph_format.space_before   = Pt(12)
        h2.paragraph_format.space_after    = Pt(2)
        h2.paragraph_format.keep_with_next = True
    except KeyError:
        pass

    # Ask Word to auto-update all field codes on open (so TOC renders immediately)
    try:
        from docx.oxml import OxmlElement as _OE
        from docx.oxml.ns import qn as _qn
        settings_elem = doc.settings.element
        update_fields = _OE("w:updateFields")
        update_fields.set(_qn("w:val"), "true")
        settings_elem.append(update_fields)
    except Exception:
        pass


# ── Logo ─────────────────────────────────────────────────────────────────────

def _add_logo_to_header(doc: Document, brand: dict, skill_dir: Path):
    """Place brand logo in the document header (right-aligned). Falls back to text."""
    logo = brand.get("logo", {})
    if not logo.get("enabled", True):
        return

    logo_path_str = logo.get("path", "")
    logo_added    = False

    if logo_path_str:
        # Try absolute path first, then relative to skill_dir
        for candidate in [Path(logo_path_str), skill_dir / logo_path_str]:
            if candidate.exists():
                width_in = brand.get("logo_rules", {}).get("docx", {}).get("width_inches", 1.0)
                for section in doc.sections:
                    header = section.header
                    hp = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
                    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    run = hp.add_run()
                    run.add_picture(str(candidate), width=Inches(width_in))
                logo_added = True
                break

    if not logo_added and logo.get("use_text_fallback_when_missing", True):
        brand_name = brand.get("brand_name", "")
        colours    = brand.get("colours", {})
        fonts      = brand.get("fonts", {})
        for section in doc.sections:
            header = section.header
            hp = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
            hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run = hp.add_run(brand_name)
            run.font.name      = fonts.get("heading", "Calibri")
            run.font.size      = Pt(9)
            run.font.bold      = True
            run.font.color.rgb = _hex_to_rgb(colours.get("primary", "#1F3864"))


# ── Cover page ───────────────────────────────────────────────────────────────

def _add_cover_page(doc: Document, brand: dict, template: dict, instance: dict):
    colours = brand.get("colours", {})
    fonts   = brand.get("fonts", {})
    primary_hex   = colours.get("primary",   "#1F3864")
    secondary_hex = colours.get("secondary", "#FF6E30")
    muted_hex     = colours.get("muted",     "#888888")

    primary   = _hex_to_rgb(primary_hex)
    secondary = _hex_to_rgb(secondary_hex)
    muted     = _hex_to_rgb(muted_hex)
    white     = _hex_to_rgb("#FFFFFF")
    heading_font = fonts.get("heading", "Calibri")
    body_font    = fonts.get("body",    "Calibri")

    # Header band — shaded background with brand name in white
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after  = Pt(10)
    _shade_paragraph(p, primary_hex.lstrip("#"))
    run = p.add_run("  " + brand.get("brand_name", "").upper())
    run.font.name      = heading_font
    run.font.size      = Pt(11)
    run.font.bold      = True
    run.font.color.rgb = white

    # Vertical breathing space
    _add_spacer(doc, Pt(48))

    # Document title
    doc_type_label = _format_section_title(
        template.get("name", template.get("document_type", "Document"))
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(4)
    run = p.add_run(doc_type_label)
    run.font.name      = heading_font
    run.font.size      = Pt(36)
    run.font.bold      = True
    run.font.color.rgb = primary

    # Accent line under title
    _add_hr(doc, colour=secondary, thickness=6)

    # Client / project subtitle
    client  = instance.get("client_name", "")
    project = instance.get("project_name", "")
    subtitle_parts = list(filter(None, [f"For {client}" if client else "", project]))
    if subtitle_parts:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after  = Pt(0)
        run = p.add_run(" · ".join(subtitle_parts))
        run.font.name      = heading_font
        run.font.size      = Pt(18)
        run.font.color.rgb = muted

    # Date
    from datetime import date
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after  = Pt(0)
    run = p.add_run(date.today().strftime("%B %d, %Y"))
    run.font.name      = body_font
    run.font.size      = Pt(11)
    run.font.color.rgb = muted


# ── Table of Contents ─────────────────────────────────────────────────────────

def _add_table_of_contents(doc: Document, brand: dict):
    """
    Insert a Table of Contents page using a Word field code.
    Word auto-updates it on open (because updateFields=true is set in settings).
    """
    colours = brand.get("colours", {})
    fonts   = brand.get("fonts", {})

    # TOC heading — use Heading 1 style so it appears in the document structure
    p = doc.add_paragraph("Contents", style="Heading 1")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _add_hr(doc, colour=_hex_to_rgb(colours.get("secondary", "#FF6E30")),
            thickness=4, space_after=Pt(8))

    # TOC field code paragraph — Word renders this as a real TOC
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(4)

    run = para.add_run()
    run.font.name      = fonts.get("body", "Calibri")
    run.font.size      = Pt(11)
    run.font.color.rgb = _hex_to_rgb(colours.get("text", "#1E1E1E"))

    # Build the field code via XML: TOC \o "1-3" \h \z \u
    fldChar_begin = OxmlElement("w:fldChar")
    fldChar_begin.set(qn("w:fldCharType"), "begin")
    fldChar_begin.set(qn("w:dirty"), "true")

    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = ' TOC \\o "1-3" \\h \\z \\u '

    fldChar_sep = OxmlElement("w:fldChar")
    fldChar_sep.set(qn("w:fldCharType"), "separate")

    # Placeholder text shown before Word updates the field
    placeholder_run = OxmlElement("w:r")
    placeholder_rPr = OxmlElement("w:rPr")
    placeholder_rStyle = OxmlElement("w:rStyle")
    placeholder_rStyle.set(qn("w:val"), "a0")  # TOCLink style
    placeholder_rPr.append(placeholder_rStyle)
    placeholder_run.append(placeholder_rPr)
    placeholder_t = OxmlElement("w:t")
    placeholder_t.text = "[ Open in Word to view Table of Contents — it updates automatically ]"
    placeholder_run.append(placeholder_t)

    fldChar_end = OxmlElement("w:fldChar")
    fldChar_end.set(qn("w:fldCharType"), "end")

    r = run._r
    r.append(fldChar_begin)
    r.append(instrText)
    r.append(fldChar_sep)
    r.append(placeholder_run)
    r.append(fldChar_end)

    doc.add_page_break()


# ── Section rendering ─────────────────────────────────────────────────────────

def _add_section(doc: Document, section_id: str, content: dict, context: dict,
                 instance: dict, brand: dict, template: dict,
                 add_page_break: bool = False):
    colours = brand.get("colours", {})
    fonts   = brand.get("fonts", {})
    spacing = brand.get("spacing", {})

    secondary_hex = colours.get("secondary", "#FF6E30")
    text_hex      = colours.get("text",      "#1E1E1E")
    body_font     = fonts.get("body",    "Calibri")
    para_gap      = spacing.get("paragraph_spacing_pt",  8)

    # Section heading — use Heading 1 style for proper document structure
    heading_text = _format_section_title(section_id)
    p = doc.add_paragraph(heading_text, style="Heading 1")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if add_page_break:
        p.paragraph_format.page_break_before = True

    _add_hr(doc, colour=_hex_to_rgb(secondary_hex), thickness=4, space_after=Pt(8))

    # Body content
    body_text = content.get(section_id)
    if not body_text:
        body_text = _derive_from_context(section_id, context, instance, template)

    if body_text:
        body_colour = _hex_to_rgb(text_hex)
        for chunk in re.split(r"\n\n+", body_text.strip()):
            chunk = chunk.strip()
            if not chunk:
                continue

            # Detect Markdown table: chunk contains at least one | separator row
            if _is_markdown_table(chunk):
                _render_table(doc, chunk, brand)
                continue

            lines     = chunk.splitlines()
            non_empty = [l for l in lines if l.strip()]
            bullet_lines = [l for l in non_empty if re.match(r"^[-•]\s|^\*\s", l.strip())]
            is_pure_bullet = len(bullet_lines) == len(non_empty) and len(non_empty) > 0

            if is_pure_bullet:
                for line in lines:
                    clean = re.sub(r"^[-•]\s+|^\*\s+", "", line.strip())
                    if not clean:
                        continue
                    bp = doc.add_paragraph(style="List Bullet")
                    bp.paragraph_format.space_after  = Pt(4)
                    bp.paragraph_format.left_indent  = Inches(0.25)
                    _render_inline(bp, clean, body_font, Pt(11), body_colour)
            else:
                for line in lines:
                    line = line.strip()
                    if not line:
                        continue
                    if re.match(r"^[-•]\s|^\*\s", line):
                        clean = re.sub(r"^[-•]\s+|^\*\s+", "", line)
                        bp = doc.add_paragraph(style="List Bullet")
                        bp.paragraph_format.space_after = Pt(4)
                        bp.paragraph_format.left_indent = Inches(0.25)
                        _render_inline(bp, clean, body_font, Pt(11), body_colour)
                    elif re.match(r"^\*\*[^*]+\*\*$", line.strip()):
                        # Entirely bold line → Heading 2 sub-heading
                        clean = line.strip().strip("*")
                        sh = doc.add_paragraph(clean, style="Heading 2")
                        sh.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    else:
                        np = doc.add_paragraph()
                        np.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        np.paragraph_format.space_after       = Pt(para_gap)
                        np.paragraph_format.space_before      = Pt(0)
                        np.paragraph_format.first_line_indent = Pt(0)
                        np.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
                        np.paragraph_format.line_spacing      = Pt(11 * 1.15)
                        _render_inline(np, line, body_font, Pt(11), body_colour)
    else:
        np = doc.add_paragraph(f"[{heading_text} — to be completed]")
        np.paragraph_format.space_after = Pt(8)
        for run in np.runs:
            run.font.name      = body_font
            run.font.size      = Pt(11)
            run.font.color.rgb = _hex_to_rgb("#AAAAAA")
            run.font.italic    = True


# ── Table rendering ───────────────────────────────────────────────────────────

def _is_markdown_table(text: str) -> bool:
    """Return True if the text block looks like a Markdown table."""
    lines = [l.strip() for l in text.strip().splitlines() if l.strip()]
    if len(lines) < 2:
        return False
    # Must have at least 2 pipe-separated lines, with a separator row (--- cells)
    pipe_lines = [l for l in lines if "|" in l]
    sep_lines  = [l for l in lines if re.match(r"^\|[\s\-|:]+\|$", l)]
    return len(pipe_lines) >= 2 and len(sep_lines) >= 1


def _render_table(doc: Document, table_text: str, brand: dict):
    """Parse a Markdown table and add a styled python-docx table."""
    colours      = brand.get("colours", {})
    fonts        = brand.get("fonts", {})
    primary_hex  = colours.get("primary",  "#1F3864").lstrip("#")
    surface_hex  = colours.get("surface",  "#F5F5F5").lstrip("#")
    text_hex     = colours.get("text",     "#1E1E1E")
    heading_font = fonts.get("heading", "Calibri")
    body_font    = fonts.get("body",    "Calibri")

    lines = [l.strip() for l in table_text.strip().splitlines() if l.strip()]
    # Split each line by | and strip cells
    rows = []
    for line in lines:
        if re.match(r"^\|[\s\-|:]+\|$", line):
            continue  # Skip separator row
        cells = [c.strip() for c in line.strip("|").split("|")]
        if cells:
            rows.append(cells)

    if not rows:
        return

    num_cols = max(len(r) for r in rows)
    table    = doc.add_table(rows=len(rows), cols=num_cols)
    table.style = "Table Grid"

    for r_idx, row_data in enumerate(rows):
        row = table.rows[r_idx]
        is_header = (r_idx == 0)
        for c_idx in range(num_cols):
            cell      = row.cells[c_idx]
            cell_text = row_data[c_idx] if c_idx < len(row_data) else ""

            # Clear default paragraph and add styled one
            para = cell.paragraphs[0]
            para.clear()
            _render_inline(para, cell_text,
                           heading_font if is_header else body_font,
                           Pt(11), _hex_to_rgb(text_hex))

            # Style runs
            for run in para.runs:
                run.font.bold = is_header

            # Cell background
            tc   = cell._tc
            tcPr = tc.get_or_add_tcPr()
            shd  = OxmlElement("w:shd")
            shd.set(qn("w:val"),   "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"),  primary_hex if is_header else
                    (surface_hex if r_idx % 2 == 0 else "FFFFFF"))
            tcPr.append(shd)

            # Header text white
            if is_header:
                for run in para.runs:
                    run.font.color.rgb = _hex_to_rgb("#FFFFFF")

    # Add spacing after table
    doc.add_paragraph().paragraph_format.space_before = Pt(8)


# ── Inline markdown renderer ──────────────────────────────────────────────────

def _render_inline(para, text: str, font_name: str, font_size: Pt, colour: RGBColor):
    """Render **bold**, *italic*, `code`, and literal \\n into a paragraph's runs."""
    parts = re.split(r"(\*\*[^*]+?\*\*|\*[^*]+?\*|`[^`]+?`)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = para.add_run(part[2:-2])
            run.font.bold   = True
            run.font.italic = False
        elif part.startswith("*") and part.endswith("*"):
            run = para.add_run(part[1:-1])
            run.font.italic = True
            run.font.bold   = False
        elif part.startswith("`") and part.endswith("`"):
            run = para.add_run(part[1:-1])
            run.font.name = "Courier New"
        else:
            segments = part.split("\n")
            for i, seg in enumerate(segments):
                run = para.add_run(seg)
                if i < len(segments) - 1:
                    run.add_break(WD_BREAK.LINE)
        last_run = para.runs[-1] if para.runs else None
        if last_run:
            if last_run.font.name not in ("Courier New",):
                last_run.font.name = font_name
            last_run.font.size      = font_size
            last_run.font.color.rgb = colour


# ── Context fallback ──────────────────────────────────────────────────────────

def _derive_from_context(section_id: str, context: dict, instance: dict, template: dict) -> str:
    """Build section content from context data when no pre-authored content exists."""

    # Pricing
    if section_id in ("pricing", "investment") and "pricing.current" in context:
        prices = context["pricing.current"].get("prices", [])
        notes  = context["pricing.current"].get("notes", "")
        lines  = [
            f"- {p.get('service_id','Service')}: "
            f"{p.get('currency','USD')} {p.get('price',0):,} ({p.get('billing','')})"
            for p in prices
        ]
        if notes:
            lines.append(f"\n{notes}")
        return "\n".join(lines)

    # Services / scope
    if section_id in ("services", "scope_of_work") and "services.catalog" in context:
        services = context["services.catalog"].get("services", [])
        lines = []
        for svc in services:
            lines.append(f"**{svc.get('name', '')}**")
            lines.append(svc.get("description", ""))
            deliverables = svc.get("deliverables", [])
            if deliverables:
                lines.append("Deliverables: " + ", ".join(deliverables))
            lines.append("")
        return "\n".join(lines)

    # Case studies — support both singular (legacy) and plural (new) schema
    if section_id == "case_studies":
        # New: proof.case_studies (array)
        if "proof.case_studies" in context:
            case_studies = context["proof.case_studies"].get("case_studies", [])
            blocks = []
            for cs in case_studies:
                blocks.append(
                    f"**{cs.get('client', '')}**\n\n"
                    f"*Challenge:* {cs.get('challenge', '')}\n\n"
                    f"*Solution:* {cs.get('solution', '')}\n\n"
                    f"*Outcome:* {cs.get('outcome', '')}"
                )
            return "\n\n---\n\n".join(blocks)
        # Legacy: proof.case_study (single object)
        if "proof.case_study" in context:
            cs = context["proof.case_study"]
            return (
                f"**{cs.get('client', '')}**\n\n"
                f"*Challenge:* {cs.get('challenge', '')}\n\n"
                f"*Solution:* {cs.get('solution', '')}\n\n"
                f"*Outcome:* {cs.get('outcome', '')}"
            )

    # Terms
    if section_id == "terms" and "legal.standard_clauses" in context:
        clauses = context["legal.standard_clauses"]
        lines = []
        if clauses.get("payment_terms"):
            lines.append(f"**Payment Terms:** {clauses['payment_terms']}")
        if clauses.get("confidentiality"):
            lines.append(f"**Confidentiality:** {clauses['confidentiality']}")
        if clauses.get("warranty_limitations"):
            lines.append(f"**Liability:** {clauses['warranty_limitations']}")
        return "\n\n".join(lines)

    return ""


# ── Footer ────────────────────────────────────────────────────────────────────

def _add_footer(doc: Document, brand: dict):
    brand_name = brand.get("brand_name", "")
    muted      = _hex_to_rgb(brand.get("colours", {}).get("muted", "#888888"))
    for section in doc.sections:
        footer = section.footer
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.clear()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"{brand_name}  ·  ")
        run.font.size      = Pt(9)
        run.font.color.rgb = muted
        _add_page_number_field(p, muted)


def _add_page_number_field(para, colour: RGBColor):
    fldChar1   = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText  = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = " PAGE "
    fldChar2   = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")
    run = para.add_run()
    run.font.size      = Pt(9)
    run.font.color.rgb = colour
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)


# ── Low-level helpers ─────────────────────────────────────────────────────────

def _add_hr(doc: Document, colour: Optional[RGBColor] = None,
            thickness: int = 6, space_after: Optional[Pt] = None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = space_after if space_after is not None else Pt(6)
    pPr    = p._p.get_or_add_pPr()
    pBdr   = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"),   "single")
    bottom.set(qn("w:sz"),    str(thickness))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), _rgb_to_hex(colour) if colour else "FF6E30")
    pBdr.append(bottom)
    pPr.append(pBdr)


def _shade_paragraph(para, hex_fill: str):
    pPr = para._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  hex_fill.lstrip("#").upper())
    pPr.append(shd)


def _add_spacer(doc: Document, height: Pt):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = height
    p.paragraph_format.space_after  = Pt(0)


def _format_section_title(section_id: str) -> str:
    labels = {
        "cover_page":           "Cover Page",
        "executive_summary":    "Executive Summary",
        "client_context":       "Client Context",
        "proposed_solution":    "Proposed Solution",
        "scope_of_work":        "Scope of Work",
        "timeline":             "Timeline",
        "pricing":              "Investment",
        "next_steps":           "Next Steps",
        "team_credentials":     "Our Team",
        "case_studies":         "Case Studies",
        "assumptions":          "Assumptions",
        "appendix":             "Appendix",
        "terms":                "Terms & Conditions",
        "problem_statement":    "Problem Statement",
        "solution_overview":    "Solution Overview",
        "features_benefits":    "Features & Benefits",
        "proof_points":         "Proof Points",
        "pricing_and_packages": "Pricing & Packages",
        "call_to_action":       "Call to Action",
        "client_background":    "Client Background",
        "challenge":            "The Challenge",
        "solution":             "Our Solution",
        "results":              "Results",
        "testimonial":          "Testimonial",
        "introduction":         "Introduction",
        "methodology":          "Methodology",
        "findings":             "Findings",
        "recommendations":      "Recommendations",
        "conclusion":           "Conclusion",
        "header":               "Report Header",
        "summary":              "Summary",
        "details":              "Details",
        "action_items":         "Action Items",
        "parties":              "Parties",
        "signatures":           "Signatures",
    }
    return labels.get(section_id, section_id.replace("_", " ").title())


def _hex_to_rgb(hex_colour: str) -> RGBColor:
    h = hex_colour.lstrip("#")
    if len(h) != 6:
        h = "1E1E1E"
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _rgb_to_hex(rgb_colour: RGBColor) -> str:
    return str(rgb_colour)
