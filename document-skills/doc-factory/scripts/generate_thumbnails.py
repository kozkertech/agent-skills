#!/usr/bin/env python3
"""
Pre-render first-page thumbnails for every template in the library.

For each template, generates a realistic sample document using the Apex Creative
Studio context library, converts it to PDF with LibreOffice, renders page 1 to a PNG,
and saves it to assets/previews/thumbnails/{template_id}.png

Run this once after initial setup, or again after adding new templates.

Usage:
    python scripts/generate_thumbnails.py [--force] [--template <id>]

Options:
    --force         Re-generate even if thumbnail already exists
    --template <id> Only generate for one specific template
    --dpi <N>       Resolution (default: 180)
"""

import sys
import json
import pathlib
import subprocess
import tempfile
import shutil

SKILL_DIR    = pathlib.Path(__file__).parent.parent
TEMPLATES_DIR = SKILL_DIR / "assets" / "templates"
CONTEXT_FILE  = SKILL_DIR / "assets" / "context" / "context_library.json"
BRAND_FILE    = SKILL_DIR / "assets" / "brand" / "brand_config.json"
THUMB_DIR     = SKILL_DIR / "assets" / "previews" / "thumbnails"
NON_TEMPLATE  = {"business_types.json"}

# Sample instance inputs per template type
SAMPLE_INPUTS = {
    "proposal": {
        "client_name": "Acme Ventures",
        "project_title": "Brand Strategy & Website Redesign",
        "scope_items": "Full brand strategy, visual identity design, responsive website build",
        "fee_total": "£24,000",
        "valid_until": "2026-04-30",
    },
    "sow": {
        "client_name": "Acme Ventures",
        "project_title": "Brand Strategy & Website Redesign",
        "scope_items": "Brand strategy, visual identity, website design and development",
        "fee_total": "£24,000",
        "start_date": "2026-04-01",
        "end_date": "2026-07-31",
    },
    "contract": {
        "client_name": "Nexus Systems Ltd",
        "project_title": "Digital Transformation Engagement",
        "fee_total": "£35,000",
        "start_date": "2026-04-01",
        "end_date": "2026-12-31",
    },
    "nda": {
        "client_name": "BrightPath Technologies",
        "purpose": "Discussions regarding a potential joint product development arrangement",
        "nda_type": "Mutual",
        "term_years": "2",
    },
    "prd": {
        "product_name": "Client Portal v2",
        "feature_name": "Real-Time Project Dashboard",
        "problem_statement": "Clients lack visibility into project progress without emailing the team",
        "target_users": "External clients with active projects",
    },
    "case-study": {
        "client_name": "TechGear Online",
        "industry": "E-commerce / Retail",
        "challenge": "Outdated website with poor conversion rates and high cart abandonment",
        "solution": "Full UX research, responsive redesign, and checkout optimisation",
        "result": "42% increase in conversion rate, 35% reduction in cart abandonment",
    },
    "sales-deck": {
        "prospect_name": "HealthCore Solutions",
        "industry": "Healthcare Technology",
        "pain_points": "Brand inconsistency across product lines, outdated digital presence",
    },
    "pitch-deck": {
        "company_name": "Apex Creative Studio",
        "tagline": "Strategic Design for Digital Growth",
        "funding_ask": "Series A — £2M",
        "use_of_funds": "Team expansion, product development, market expansion",
    },
    "price-list": {
        "currency": "GBP",
        "valid_date": "January–December 2026",
    },
    "status-update": {
        "project_name": "Horizon Platform",
        "reporting_period": "Week of March 17, 2026",
        "rag_status": "Amber",
        "progress": "Completed UI design for all 5 screens",
        "plans": "Complete API integration layer",
        "problems": "Third-party API documentation is outdated, risking a 1-week delay",
    },
    "runbook": {
        "process_name": "Monthly Client Report Generation",
        "system": "Client Portal",
        "owner": "Sarah Chen",
    },
    "technical-report": {
        "project_name": "E-commerce Platform Migration",
        "client_name": "Retail Solutions Inc",
        "report_type": "Technical Architecture Review",
    },
    "portfolio": {
        "portfolio_title": "Selected Work 2024–2025",
        "focus": "Brand Strategy & Digital Design",
    },
}


def load_json(path):
    try:
        return json.loads(path.read_text())
    except Exception:
        return {}


def build_docx_sample(template_id, template, context, brand, inputs):
    """Generate a sample DOCX for the given template using python-docx."""
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor, Inches, Cm
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
    except ImportError:
        return None

    doc = Document()
    company = context.get("company_profile", {})
    brand_hex = brand.get("colors", {}).get("primary", "#1F3A70").lstrip("#")
    brand_rgb = tuple(int(brand_hex[i:i+2], 16) for i in (0, 2, 4))
    font_name = brand.get("fonts", {}).get("body", "Arial")
    brand_color = RGBColor(*brand_rgb)

    # Set page margins
    from docx.oxml.ns import nsmap
    for section in doc.sections:
        section.top_margin    = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin   = Cm(3.0)
        section.right_margin  = Cm(3.0)

    # Cover block
    cover = doc.add_paragraph()
    cover.paragraph_format.space_before = Pt(0)

    # Company name (styled)
    p = doc.add_paragraph()
    run = p.add_run(company.get("name", "Apex Creative Studio"))
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = brand_color
    run.font.name = font_name
    p.paragraph_format.space_after = Pt(4)

    # Tagline
    p2 = doc.add_paragraph()
    run2 = p2.add_run(company.get("tagline", "Strategic Design for Digital Growth"))
    run2.font.size = Pt(12)
    run2.font.color.rgb = RGBColor(100, 116, 139)
    run2.font.name = font_name
    p2.paragraph_format.space_after = Pt(20)

    # Document title
    title_text = template.get("name", template_id.replace("-", " ").title())
    client = inputs.get("client_name", "")
    if client:
        title_text = f"{title_text}\n{client}"

    title_p = doc.add_paragraph()
    title_run = title_p.add_run(template.get("name", title_text))
    title_run.bold = True
    title_run.font.size = Pt(22)
    title_run.font.name = font_name
    title_p.paragraph_format.space_after = Pt(6)

    if client:
        sub_p = doc.add_paragraph()
        sub_run = sub_p.add_run(f"Prepared for: {client}")
        sub_run.font.size = Pt(13)
        sub_run.font.color.rgb = RGBColor(71, 85, 105)
        sub_run.font.name = font_name
        sub_p.paragraph_format.space_after = Pt(4)

    # Project/fee line
    project = inputs.get("project_title", "")
    fee = inputs.get("fee_total", "")
    if project or fee:
        detail_p = doc.add_paragraph()
        detail_text = " | ".join(filter(None, [project, fee]))
        detail_run = detail_p.add_run(detail_text)
        detail_run.font.size = Pt(11)
        detail_run.font.color.rgb = brand_color
        detail_run.font.name = font_name
        detail_p.paragraph_format.space_after = Pt(20)

    # Horizontal rule
    hr = doc.add_paragraph()
    p_pr = hr._p.get_or_add_pPr()
    p_bdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), brand_hex)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)
    hr.paragraph_format.space_after = Pt(16)

    # Sections
    sections_data = template.get("required_sections", [])[:8]
    for sec in sections_data:
        # Section heading
        h = doc.add_paragraph()
        h_run = h.add_run(sec.get("title", sec.get("id", "Section").replace("_", " ").title()))
        h_run.bold = True
        h_run.font.size = Pt(14)
        h_run.font.color.rgb = brand_color
        h_run.font.name = font_name
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after  = Pt(6)

        # Section body (placeholder text with description)
        desc = sec.get("description", "")
        body_text = desc if desc else f"[{sec.get('title', 'Content')}]"
        b = doc.add_paragraph()
        b_run = b.add_run(body_text)
        b_run.font.size = Pt(10.5)
        b_run.font.color.rgb = RGBColor(71, 85, 105)
        b_run.font.name = font_name
        b.paragraph_format.space_after = Pt(8)

        # Add some filler paragraphs to make it feel like a real document
        for _ in range(2):
            filler = doc.add_paragraph()
            filler_run = filler.add_run(
                "Lorem ipsum dolor sit amet, consectetur adipiscing elit. "
                "Sed do eiusmod tempor incididunt ut labore et dolore magna aliqua. "
                "Ut enim ad minim veniam, quis nostrud exercitation."
            )
            filler_run.font.size = Pt(10.5)
            filler_run.font.color.rgb = RGBColor(100, 116, 139)
            filler_run.font.name = font_name
            filler.paragraph_format.space_after = Pt(4)

    # Footer info
    footer_p = doc.add_paragraph()
    footer_p.paragraph_format.space_before = Pt(20)
    footer_run = footer_p.add_run(
        f"{company.get('name','')}  ·  {company.get('website','')}  ·  {company.get('email','')}"
    )
    footer_run.font.size = Pt(9)
    footer_run.font.color.rgb = RGBColor(148, 163, 184)
    footer_run.font.name = font_name
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    return doc


def build_xlsx_sample(template_id, context, brand):
    """Generate a sample XLSX for price-list template."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        return None

    brand_hex = brand.get("colors", {}).get("primary", "#1F3A70").lstrip("#")
    company = context.get("company_profile", {})
    services = context.get("services", [])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Price List"

    hdr_fill = PatternFill("solid", start_color=brand_hex, end_color=brand_hex)
    hdr_font = Font(bold=True, color="FFFFFF", name="Arial", size=11)
    alt_fill = PatternFill("solid", start_color="F9FAFB", end_color="F9FAFB")
    total_fill = PatternFill("solid", start_color="EFF6FF", end_color="EFF6FF")
    brand_font = Font(bold=True, color=brand_hex, name="Arial", size=10)
    body_font = Font(name="Arial", size=10)

    thin = Side(style="thin", color="E2E8F0")
    border = Border(bottom=thin)

    # Title row
    ws.merge_cells("A1:D1")
    ws["A1"] = f"{company.get('name','Price List')} — Service Pricing 2026"
    ws["A1"].font = Font(bold=True, name="Arial", size=14, color=brand_hex)
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.row_dimensions[2].height = 8  # spacer

    # Header
    headers = ["Service", "Description", "Pricing", "Notes"]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=c, value=h)
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[3].height = 22

    # Data rows
    for i, svc in enumerate(services[:8], start=4):
        bg = alt_fill if i % 2 == 0 else PatternFill("solid", start_color="FFFFFF", end_color="FFFFFF")
        row = [
            svc.get("name", ""),
            svc.get("description", ""),
            svc.get("pricing", ""),
            svc.get("unit", ""),
        ]
        for c, val in enumerate(row, start=1):
            cell = ws.cell(row=i, column=c, value=val)
            cell.fill = bg
            cell.font = body_font
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = border
        ws.row_dimensions[i].height = 36

    # Column widths
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 45
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 18

    return wb


def render_thumbnail(file_path: pathlib.Path, out_png: pathlib.Path, dpi: int = 180):
    """Convert a document to PDF and extract page 1 as a PNG thumbnail."""
    with tempfile.TemporaryDirectory() as tmpdir:
        work = pathlib.Path(tmpdir)

        # Convert to PDF
        result = subprocess.run(
            ["libreoffice", "--headless", "--convert-to", "pdf",
             "--outdir", str(work), str(file_path)],
            capture_output=True, text=True, timeout=60
        )
        if result.returncode != 0:
            print(f"    ⚠ LibreOffice failed: {result.stderr[:200]}")
            return False

        pdfs = list(work.glob("*.pdf"))
        if not pdfs:
            print(f"    ⚠ No PDF produced")
            return False

        # Render page 1 only
        result2 = subprocess.run(
            ["pdftoppm", "-r", str(dpi), "-png", "-l", "1",
             str(pdfs[0]), str(work / "page")],
            capture_output=True, text=True, timeout=30
        )
        if result2.returncode != 0:
            print(f"    ⚠ pdftoppm failed: {result2.stderr[:200]}")
            return False

        pages = sorted(work.glob("page-*.png")) or sorted(work.glob("page*.png"))
        if not pages:
            print(f"    ⚠ No page PNG produced")
            return False

        out_png.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy(pages[0], out_png)
        return True


def main():
    args = sys.argv[1:]
    force = "--force" in args
    target = None
    dpi = 180
    if "--template" in args:
        i = args.index("--template")
        target = args[i + 1]
    if "--dpi" in args:
        i = args.index("--dpi")
        dpi = int(args[i + 1])

    THUMB_DIR.mkdir(parents=True, exist_ok=True)
    context = load_json(CONTEXT_FILE)
    brand   = load_json(BRAND_FILE)

    templates = {}
    for f in sorted(TEMPLATES_DIR.glob("*.json")):
        if f.name in NON_TEMPLATE:
            continue
        try:
            data = json.loads(f.read_text())
            if "id" in data:
                templates[data["id"]] = data
        except Exception:
            continue

    if target:
        templates = {k: v for k, v in templates.items() if k == target}
        if not templates:
            print(f"Template '{target}' not found.")
            sys.exit(1)

    print(f"Generating thumbnails for {len(templates)} template(s) at {dpi} dpi…\n")

    success = 0
    for tid, template in templates.items():
        out_png = THUMB_DIR / f"{tid}.png"
        if out_png.exists() and not force:
            print(f"  ✓ {tid} (cached)")
            success += 1
            continue

        print(f"  → {tid}…", end=" ", flush=True)
        inputs = SAMPLE_INPUTS.get(tid, {})
        fmt = template.get("format", "docx")

        try:
            with tempfile.TemporaryDirectory() as tmpdir:
                tmp = pathlib.Path(tmpdir)
                if fmt == "xlsx":
                    wb = build_xlsx_sample(tid, context, brand)
                    if wb is None:
                        print("skipped (openpyxl missing)")
                        continue
                    tmp_file = tmp / f"{tid}.xlsx"
                    wb.save(tmp_file)
                else:
                    doc = build_docx_sample(tid, template, context, brand, inputs)
                    if doc is None:
                        print("skipped (python-docx missing)")
                        continue
                    tmp_file = tmp / f"{tid}.docx"
                    doc.save(tmp_file)

                ok = render_thumbnail(tmp_file, out_png, dpi=dpi)
                if ok:
                    kb = out_png.stat().st_size // 1024
                    print(f"✓ ({kb} KB)")
                    success += 1
                else:
                    print("failed")

        except Exception as e:
            print(f"error: {e}")

    print(f"\nDone: {success}/{len(templates)} thumbnails in {THUMB_DIR}")


if __name__ == "__main__":
    main()
